import customtkinter as ctk
from tkinter import filedialog
from docx import Document
from fpdf import FPDF
from fpdf.enums import XPos, YPos
from frames import *
from widgets import *
from settings import *

class App(ctk.CTk):

    def __init__(self):

        super().__init__()
        self.title('Ghosting Text')
        self.minsize(W_WIDTH, W_HEIGHT)
        self.maxsize(W_WIDTH, W_HEIGHT)

        self.is_on = True
        self.timer = Timer(self)
        self.typing = TypingFrame(self)
        self.instructions = InstructionsButton(self, self.instru_info)
        self.file_panel = FileTypePanel(self)
        self.export_button = ExportButton(self, self.export)
        self.browse_button = BrowseButton(self, self.path)
        self.reset_button = ResetButton(self, self.reset)

        self.filepath = None
        self.default_filename = 'my_text'

        self.after(100, lambda: self.typing.text_box.focus())
        self.bind_keys()

        self.mainloop()


    def instru_info(self):

        info_window = ctk.CTkToplevel(self)
        info_window.title("Information")
        info_window.geometry("300x200")
        label = ctk.CTkLabel(info_window,
                             text="Just type away! But careful — if you stop typing for more than 5 seconds, your work will be erased. "
                                  "Keep going until you reach the goal to save your progress. Have fun!",
                             wraplength=250)
        label.pack(pady=20)
        close_button = ctk.CTkButton(info_window, text="Cerrar", command=info_window.destroy)
        close_button.pack(pady=10)


    def reset(self):

        self.file_panel.disable_frame()
        self.browse_button.disable()
        self.export_button.disable()
        self.timer.timer_running = False
        self.timer.destroy()
        self.timer = Timer(self)
        self.typing.time_label.configure(text_color=T_COLOR)
        self.timer.color = 0
        self.typing.clean_box()
        self.typing.words_counter.configure(text=f'Words: ---/{GOAL}')
        self.typing.time_label.configure(text='Waiting')


    def path(self):

        self.filepath = filedialog.askdirectory()
        self.export_button.enable()


    def export(self):

        if self.filepath and self.file_panel.name_string.get():

            if self.file_panel.txt_var.get() == 'txt':
                with open(f"{self.filepath}/{self.file_panel.name_string.get()}.txt", "w", encoding="utf-8") as f:
                    f.write(self.typing.text_box.get("1.0", "end-1c"))
                    self.saved_success(self.filepath)

            elif self.file_panel.doc_var.get() == 'doc':
                document = Document()
                document.add_paragraph(self.typing.text_box.get("1.0", "end-1c"))
                document.save(f"{self.filepath}/{self.file_panel.name_string.get()}.docx")
                self.saved_success(self.filepath)

            elif self.file_panel.pdf_var.get() == 'pdf':

                pdf = FPDF()
                pdf.add_page()
                pdf.set_font("Helvetica", size=12)

                texto = self.typing.text_box.get("1.0", "end-1c")
                for linea in texto.split("\n"):
                    pdf.cell(200, 10, text=linea, new_x=XPos.LMARGIN, new_y=YPos.NEXT)

                pdf.output(f"{self.filepath}/{self.file_panel.name_string.get()}.pdf")
                self.saved_success(self.filepath)

            else:
                self.saved_error()

        else:
            self.saved_error()


    def saved_success(self, filepath):

        success_label = ctk.CTkLabel(self, text=f"Text saved in:\n{filepath}", fg_color="green", text_color="white", corner_radius=5)
        success_label.place(relx=0.5, rely=0.9, anchor='center')
        self.after(3000, success_label.destroy)  # 3 secs


    def saved_error(self):

        error_label = ctk.CTkLabel(self, text=f"Error saving, something is missing.", fg_color="red", text_color="white", corner_radius=5)
        error_label.place(relx=0.5, rely=0.9, anchor='center')
        self.after(5000, error_label.destroy)


    def time_over(self):

        self.typing.clean_box()


    def bind_keys(self):

        self.typing.text_box.bind('<Key>', self.key_detected)


    def key_detected(self, event):

        if event.keysym == 'space':
            self.typing.count_words()
            if self.typing.check_goal():
                self.is_on = False
                self.timer.timer_running = False
                self.typing.time_label.configure(text='Congratulations 🎉, You did It! Now you can export you text 👍🏼', text_color=T_GREEN)
                self.browse_button.enable()
                self.file_panel.enable_frame()
        elif event.keysym == 'BackSpace':
            return
        if self.is_on:
            if not self.timer.timer_running:
                self.timer.reset()
                self.timer.timer_running = True
                self.timer.start(self.time_over,self.typing.time_label, self.typing.text_box)
                return
            else:
                self.timer.timer_running = False
                self.timer.reset()
                self.timer.destroy()
                self.timer = Timer(self)
                self.typing.time_label.configure(text_color=T_COLOR)
                self.timer.color = 0
                self.timer.timer_running = True
                self.timer.start(self.time_over, self.typing.time_label, self.typing.text_box)
                return


App()